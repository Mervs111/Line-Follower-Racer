"""RAB CAPEX & OPEX pemeliharaan 5 tahun - armada AGV mini line tracer (5 unit). Semua angka pakai rumus.
Pakai: python3 gen_rab_pemeliharaan.py RAB.xlsx   (butuh openpyxl; Excel menghitung rumus saat file dibuka)"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.page import PageMargins

OUT = sys.argv[1]
FONT = "Arial"
BLUE, BLACK, GREEN, GREY, WHITE = "0000FF", "000000", "008000", "808080", "FFFFFF"
NAVY = "1F2C57"
F_IN = PatternFill("solid", fgColor="FFFF00")        # asumsi kunci (boleh diubah)
F_HEAD = PatternFill("solid", fgColor=NAVY)
F_SEC = PatternFill("solid", fgColor="DCE3F0")
F_TOT = PatternFill("solid", fgColor="EEF1F6")
THIN = Side(style="thin", color="B8C2D3")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
RP = '#,##0;(#,##0);"-"'
RP2 = '#,##0.00;(#,##0.00);"-"'
NUM = '#,##0;(#,##0);"-"'
NUM1 = '#,##0.0;(#,##0.0);"-"'
PCT = '0.0%;(0.0%);"-"'

wb = Workbook()


def font(color=BLACK, bold=False, size=10, italic=False):
    return Font(name=FONT, color=color, bold=bold, size=size, italic=italic)


def put(ws, ref, value, color=BLACK, bold=False, fmt=None, fill=None, wrap=False, align=None, size=10, italic=False,
        border=True):
    c = ws[ref]
    c.value = value
    c.font = font(color, bold, size, italic)
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    c.alignment = Alignment(wrap_text=wrap, vertical="top", horizontal=align)
    if border:
        c.border = BOX
    return c


def header(ws, row, labels, widths=None, height=30):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=i + 1, value=lab)
        c.font = font(WHITE, True)
        c.fill = F_HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = BOX
    ws.row_dimensions[row].height = height
    if widths:
        for i, w in enumerate(widths):
            ws.column_dimensions[get_column_letter(i + 1)].width = w


def title(ws, text, sub=None, sub2=None):
    put(ws, "A1", text, bold=True, size=14, border=False)
    if sub:
        put(ws, "A2", sub, color=GREY, border=False, italic=True)
    if sub2:
        put(ws, "A3", sub2, color=GREY, border=False, italic=True)


def name(nm, ref):
    wb.defined_names[nm] = DefinedName(nm, attr_text=ref)


def page(ws, landscape=True):
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.5)


# ======================================================================= ASUMSI
A = wb.active
A.title = "Asumsi"
title(A, "Asumsi & Parameter Perhitungan",
      "Sel biru berlatar kuning = asumsi yang boleh diubah. Semua angka di sheet lain dihitung ulang otomatis.",
      "Sel hitam = rumus. Sel hijau = mengambil angka dari sheet lain.")
header(A, 5, ["No", "Parameter", "Nilai", "Satuan", "Sumber / dasar"], [5, 44, 16, 18, 92])
rows = [
    ("A", "Sistem & pola operasi"),
    ("n_robot", "Jumlah robot terpasang", 5, "unit", "Skenario: armada AGV mini line tracer (hasil PjBL) pengantar part/kanban antar stasiun kerja; 4 unit jalan + 1 unit cadangan rotasi", NUM),
    ("n_operasi", "Robot beroperasi bersamaan", 4, "unit", "Skenario: 1 unit selalu cadangan (dirotasi saat perawatan/pengisian)", NUM),
    ("jam_hari", "Jam operasi per hari", 8, "jam", "Skenario: 1 shift", NUM),
    ("hari_tahun", "Hari operasi per tahun", 300, "hari", "Skenario: 6 hari kerja per minggu dikurangi libur nasional", NUM),
    ("utilisasi", "Utilisasi robot selama shift", 0.6, "-", "Skenario: robot hanya berjalan saat ada misi antar (60% waktu shift)", PCT),
    ("jam_jalan", "Jam jalan armada per tahun", "=n_operasi*jam_hari*hari_tahun*utilisasi", "jam/tahun", "Dihitung: robot beroperasi x jam/hari x hari/tahun x utilisasi", NUM),
    ("jam_robot", "Rata-rata jam jalan per robot per tahun", "=jam_jalan/n_robot", "jam/tahun", "Dihitung: beban dibagi rata ke semua robot karena rotasi", NUM),
    ("B", "Umur & keandalan komponen"),
    ("runtime", "Waktu jalan per pengisian baterai", 3, "jam", "Charter proyek: baterai 2x18650, pemakaian +/-3 jam per pengisian", NUM1),
    ("n_cas", "Jumlah pengisian baterai per tahun", "=jam_jalan/runtime", "kali/tahun", "Dihitung: jam jalan armada / waktu jalan per pengisian", NUM),
    ("umur_baterai", "Umur siklus sel 18650", 500, "siklus", "Estimasi umum sel Li-ion konsumen (300-500 siklus sampai kapasitas 80%); dipakai nilai atas karena arus beban kecil", NUM),
    ("umur_motor", "Umur motor N20 (motor DC sikat)", 1000, "jam", "Estimasi kelompok untuk motor DC sikat ukuran kecil; diverifikasi dari data riwayat tahun pertama", NUM),
    ("umur_ban", "Umur ban roda silikon", 1500, "jam", "Estimasi kelompok (lantai pabrik halus); diverifikasi dari data riwayat", NUM),
    ("umur_caster", "Umur ball caster", 3000, "jam", "Estimasi kelompok", NUM),
    ("laju_modul", "Laju kerusakan modul elektronik", 0.10, "per unit per tahun", "Estimasi kelompok: 1 dari 10 modul rusak per tahun (salah colok, benturan, ESD)", PCT),
    ("laju_sensor", "Laju kerusakan sensor IR", 0.05, "per sensor per tahun", "Estimasi kelompok: sensor di bawah PCB jarang terbentur", PCT),
    ("mttr", "Waktu rata-rata penggantian komponen (MTTR)", 1, "jam per kejadian", "Estimasi: modul dipasang lewat header/konektor; termasuk uji ulang di lintasan", NUM1),
    ("C", "Tenaga kerja"),
    ("umk", "UMK Kota Surabaya 2026", 5288796, "Rp/bulan", "Kepgub Jatim No. 100.3.3.1/937/013/2025 (berlaku 1 Jan 2026) - https://www.jawapos.com/surabaya-raya/2512250021/resmi-ini-daftar-umk-38-kabupatenkota-di-jawa-timur-2026-surabaya-nyaris-rp-53-juta", RP),
    ("beban_upah", "Beban BPJS + THR pemberi kerja", 0.189, "dari upah", "BPJS Kesehatan 4% + BPJS Ketenagakerjaan +/-6,5% (JHT 3,7%, JP 2%, JKK 0,54%, JKM 0,3%) + THR 1/12 (8,33%)", PCT),
    ("jam_bulan", "Pembagi upah per jam", 173, "jam/bulan", "PP 35/2021: upah sejam = 1/173 x upah sebulan", NUM),
    ("tarif_teknisi", "Biaya teknisi per jam", "=umk*(1+beban_upah)/jam_bulan", "Rp/jam", "Dihitung: UMK x (1 + beban) / 173. Teknisi pemeliharaan diasumsikan dibayar setara UMK Surabaya", RP),
    ("D", "Ekonomi, utilitas & lintasan"),
    ("eskalasi", "Eskalasi harga per tahun", 0.035, "per tahun", "Batas atas sasaran inflasi BI 2026-2027 (2,5% +/- 1%); inflasi Juni 2026 = 3,34% - https://ekonomi.republika.co.id/berita/t8fpk6423/bi-optimistis-inflasi-20262027-tetap-terkendali-dalam-target", PCT),
    ("kontingensi", "Kontingensi OPEX", 0.10, "dari subtotal OPEX", "Cadangan untuk kerusakan di luar perkiraan dan kenaikan harga komponen impor", PCT),
    ("tarif_listrik", "Tarif listrik industri I-3/TM", 1114.74, "Rp/kWh", "Tarif PLN nonsubsidi triwulan IV 2026 (tidak berubah dari triwulan III) - https://www.fortuneidn.com/news/tarif-listrik-per-1-oktober-2026-d7w02-00-lflgw-zwfyxs", RP2),
    ("v_baterai", "Tegangan nominal baterai 2S", 7.4, "V", "Spesifikasi robot (2x18650 seri)", NUM1),
    ("kap_baterai", "Kapasitas sel 18650", 2.6, "Ah", "Sel 2.600 mAh", NUM1),
    ("efisiensi", "Efisiensi pengisian", 0.85, "-", "Estimasi charger boost 2S USB-C", PCT),
    ("wh_cas", "Energi listrik per pengisian", "=v_baterai*kap_baterai/efisiensi", "Wh", "Dihitung: V x Ah / efisiensi", NUM1),
    ("panjang_lintasan", "Panjang garis lintasan di lantai produksi", 60, "m", "Skenario tata letak: rute antar 4 stasiun dengan persimpangan 90 derajat", NUM),
    ("ganti_lintasan", "Bagian garis yang diganti per tahun", 0.5, "dari panjang", "Estimasi: pita aus karena lalu lintas troli dan pejalan kaki", PCT),
    ("pagu", "Pagu anggaran 5 tahun", 100000000, "Rp", "Ketentuan tugas: dana maksimal Rp 100.000.000", RP),
    ("tahun_awal", "Tahun anggaran pertama", 2027, "-", "Asumsi: sistem sudah beroperasi, anggaran pemeliharaan mulai tahun anggaran 2027", "0"),
]
r = 6
no = 0
for item in rows:
    if len(item) == 2:
        code, lab = item
        put(A, f"A{r}", code, bold=True, fill=F_SEC)
        put(A, f"B{r}", lab, bold=True, fill=F_SEC)
        for col in "CDE":
            put(A, f"{col}{r}", None, fill=F_SEC)
        r += 1
        continue
    nm, lab, val, unit, src, fmt = item
    no += 1
    put(A, f"A{r}", no, align="center")
    put(A, f"B{r}", lab, wrap=True)
    is_input = not (isinstance(val, str) and val.startswith("="))
    put(A, f"C{r}", val, color=BLUE if is_input else BLACK, fmt=fmt, fill=F_IN if is_input else None)
    put(A, f"D{r}", unit, wrap=True)
    put(A, f"E{r}", src, wrap=True)
    name(nm, f"Asumsi!$C${r}")
    r += 1
A.freeze_panes = "A6"
page(A)

# ======================================================================= JADWAL PM
J = wb.create_sheet("Jadwal PM")
title(J, "Jadwal Pemeliharaan Preventif (per tahun)",
      "Jam kerja teknisi dari tabel ini menjadi biaya tenaga kerja preventif di sheet OPEX.",
      "Pemeriksaan awal shift dikerjakan operator (autonomous maintenance), jamnya dicatat tapi tidak dibiayakan di RAB ini.")
header(J, 5, ["No", "Kegiatan", "Rincian pekerjaan", "Pelaksana", "Frekuensi", "Kali per tahun",
              "Durasi per objek (menit)", "Dasar", "Jumlah objek", "Jam per tahun", "Alat utama"],
       [5, 30, 62, 11, 12, 10, 11, 11, 9, 10, 34])
pm = [
    ("Pemeriksaan awal shift", "Cek tegangan baterai di OLED, bersihkan sensor dengan kuas antistatis, cek roda & ball caster, uji jalan 1 putaran", "Operator", "Harian", "=hari_tahun", 3, "per robot", "Kuas antistatis"),
    ("Pemeliharaan mingguan", "Bersihkan sensor IR dengan isopropyl alcohol, kencangkan baut & roda, cek kabel/konektor, cek tegangan sel & BMS, uji di lintasan kalibrasi", "Teknisi", "Mingguan", 52, 10, "per robot", "Multimeter, set perkakas, lintasan uji"),
    ("Pemeliharaan bulanan", "Uji kapasitas & tahanan dalam baterai, ukur RPM & arus motor, cek sinyal encoder, unduh log kejadian & cek versi firmware", "Teknisi", "Bulanan", 12, 30, "per robot", "Battery analyzer, tachometer, osiloskop, laptop"),
    ("Overhaul tahunan", "Bongkar total, bersihkan, ganti komponen sesuai kondisi, kalibrasi ulang sensor & PID, uji fungsi lengkap di lintasan 4x4", "Teknisi", "Tahunan", 1, 120, "per robot", "Semua alat kerja"),
    ("Inspeksi garis lintasan", "Cek pita garis terkelupas/kotor, tambal, bersihkan area persimpangan", "Teknisi", "Mingguan", 52, 15, "per sistem", "Pita garis, cutter"),
    ("Perawatan perangkat lunak", "Backup firmware & parameter, pembaruan firmware, tinjau log kejadian & data downtime", "Teknisi", "Bulanan", 12, 60, "per sistem", "Laptop diagnostik"),
    ("Stasiun pengisian & keselamatan", "Cek adaptor & kabel pengisian, kotak penyimpanan baterai, tekanan APAR", "Teknisi", "Bulanan", 12, 30, "per sistem", "Multimeter"),
]
r0 = 6
for i, (keg, rinci, pel, frek, kali, dur, dasar, alat) in enumerate(pm):
    r = r0 + i
    put(J, f"A{r}", i + 1, align="center")
    put(J, f"B{r}", keg, wrap=True)
    put(J, f"C{r}", rinci, wrap=True)
    put(J, f"D{r}", pel)
    put(J, f"E{r}", frek)
    put(J, f"F{r}", kali, color=GREEN if str(kali).startswith("=") else BLUE, fmt=NUM)
    put(J, f"G{r}", dur, color=BLUE, fmt=NUM)
    put(J, f"H{r}", dasar)
    put(J, f"I{r}", f'=IF(H{r}="per robot",n_robot,1)', color=GREEN, fmt=NUM)
    put(J, f"J{r}", f"=F{r}*G{r}/60*I{r}", fmt=NUM1)
    put(J, f"K{r}", alat, wrap=True)
r1 = r0 + len(pm) - 1
rt = r1 + 2
put(J, f"B{rt}", "Jam teknisi per tahun (dibiayakan)", bold=True, fill=F_TOT)
put(J, f"J{rt}", f'=SUMIF(D{r0}:D{r1},"Teknisi",J{r0}:J{r1})', bold=True, fmt=NUM1, fill=F_TOT)
put(J, f"B{rt + 1}", "Jam operator per tahun (tidak dibiayakan)", fill=F_TOT)
put(J, f"J{rt + 1}", f'=SUMIF(D{r0}:D{r1},"Operator",J{r0}:J{r1})', fmt=NUM1, fill=F_TOT)
name("jam_pm", f"'Jadwal PM'!$J${rt}")
J.freeze_panes = "A6"
page(J)

# ======================================================================= SUKU CADANG
S = wb.create_sheet("Suku Cadang")
title(S, "Kebutuhan Suku Cadang per Tahun & Stok Awal Kritis",
      "Harga satuan = RAB PjBL revisi 3.4 (charter proyek). Kebutuhan dihitung dari jam jalan, siklus baterai, atau laju kerusakan di sheet Asumsi.",
      "Stok awal kritis dibeli sekali di tahun 1 (masuk CAPEX); kebutuhan per tahun masuk OPEX.")
header(S, 5, ["No", "Komponen", "Satuan", "Harga satuan (Rp)", "Jumlah per robot", "Dasar penggantian",
              "Umur (jam / siklus)", "Laju kerusakan per tahun", "Kebutuhan per tahun", "Unit per penggantian",
              "Penggantian per tahun", "Stok awal kritis", "Biaya per tahun (Rp)", "Biaya stok awal (Rp)", "Sumber harga"],
       [5, 34, 8, 12, 9, 15, 11, 11, 11, 11, 11, 9, 13, 13, 22])
SRC = "RAB PjBL rev 3.4"
parts = [
    ("Baterai Li-ion 18650 (2.600 mAh)", "pcs", 40000, 2, "Siklus baterai", "=umur_baterai", None, 2, 4, SRC),
    ("Motor N20 6V 500 RPM + encoder", "pcs", 55000, 2, "Aus (jam)", "=umur_motor", None, 1, 4, SRC),
    ("Roda + ban silikon N20", "pasang", 25000, 1, "Aus (jam)", "=umur_ban", None, 1, 2, SRC),
    ("Ball caster mini", "pcs", 10000, 1, "Aus (jam)", "=umur_caster", None, 1, 1, SRC),
    ("Sensor IR ITR20001", "pcs", 3000, 5, "Kerusakan acak", None, "=laju_sensor", 1, 5, SRC),
    ("PCB sensor jangkar (fabrikasi)", "pcs", 20000, 1, "Kerusakan acak", None, 0.10, 1, 1, "RAB PjBL rev 3.4 (paket PCB / 5 lembar)"),
    ("ESP32 DevKit 30 pin", "pcs", 76000, 1, "Kerusakan acak", None, "=laju_modul", 1, 1, SRC),
    ("Modul driver motor TB6612FNG", "pcs", 25000, 1, "Kerusakan acak", None, "=laju_modul", 1, 1, SRC),
    ("Modul buck MP1584 5 V", "pcs", 10000, 1, "Kerusakan acak", None, "=laju_modul", 1, 1, SRC),
    ("Modul BMS 2S", "pcs", 10000, 1, "Kerusakan acak", None, "=laju_modul", 1, 1, SRC),
    ("Modul charger 2S USB-C", "pcs", 25000, 1, "Kerusakan acak", None, 0.20, 1, 1, "RAB PjBL rev 3.4 (port dicolok tiap hari: laju 20%)"),
    ("OLED 1,3 inci I2C", "pcs", 52000, 1, "Kerusakan acak", None, 0.05, 1, 1, SRC),
    ("Holder baterai 2x18650", "pcs", 8000, 1, "Kerusakan acak", None, "=laju_modul", 1, 1, SRC),
    ("Saklar geser on/off", "pcs", 3000, 1, "Kerusakan acak", None, "=laju_modul", 1, 2, SRC),
    ("Tombol tekan", "pcs", 1000, 2, "Kerusakan acak", None, "=laju_modul", 1, 2, SRC),
    ("Buzzer", "pcs", 3000, 1, "Kerusakan acak", None, 0.05, 1, 1, SRC),
    ("Konektor JST PH 6 pin + kabel motor", "set", 5000, 2, "Kerusakan acak", None, 0.20, 1, 4, SRC),
    ("Paket komponen pasif (R, C, dioda SS14, header)", "paket", 25000, 1, "Tetap per tahun", None, None, 1, 1, SRC),
]
s0 = 6
for i, (nm_, sat, harga, per, mode, umur, laju, upe, stok, src) in enumerate(parts):
    r = s0 + i
    put(S, f"A{r}", i + 1, align="center")
    put(S, f"B{r}", nm_, wrap=True)
    put(S, f"C{r}", sat)
    put(S, f"D{r}", harga, color=BLUE, fmt=RP)
    put(S, f"E{r}", per, color=BLUE, fmt=NUM)
    put(S, f"F{r}", mode)
    put(S, f"G{r}", umur, color=GREEN if umur else BLACK, fmt=NUM)
    put(S, f"H{r}", laju, color=(GREEN if isinstance(laju, str) else BLUE) if laju is not None else BLACK, fmt=PCT)
    put(S, f"I{r}", (f'=IF(F{r}="Siklus baterai",ROUNDUP(n_cas/G{r}*E{r},0),'
                     f'IF(F{r}="Aus (jam)",ROUNDUP(jam_jalan*E{r}/G{r},0),'
                     f'IF(F{r}="Kerusakan acak",ROUNDUP(n_robot*E{r}*H{r},0),E{r})))'), fmt=NUM)
    put(S, f"J{r}", upe, color=BLUE, fmt=NUM)
    put(S, f"K{r}", f"=ROUNDUP(I{r}/J{r},0)", fmt=NUM)
    put(S, f"L{r}", stok, color=BLUE, fmt=NUM)
    put(S, f"M{r}", f"=I{r}*D{r}", fmt=RP)
    put(S, f"N{r}", f"=L{r}*D{r}", fmt=RP)
    put(S, f"O{r}", src, wrap=True)
s1 = s0 + len(parts) - 1
st = s1 + 2
put(S, f"B{st}", "Total", bold=True, fill=F_TOT)
for col in "KMN":
    put(S, f"{col}{st}", f"=SUM({col}{s0}:{col}{s1})", bold=True, fmt=RP if col != "K" else NUM, fill=F_TOT)
put(S, f"B{st + 1}", "Jam teknisi untuk penggantian komponen per tahun", fill=F_TOT)
put(S, f"K{st + 1}", f"=K{st}*mttr", fmt=NUM1, fill=F_TOT)
put(S, f"L{st + 1}", "jam", fill=F_TOT)
name("biaya_sc_tahun", f"'Suku Cadang'!$M${st}")
name("biaya_stok_awal", f"'Suku Cadang'!$N${st}")
name("jam_ganti", f"'Suku Cadang'!$K${st + 1}")
name("n_ganti", f"'Suku Cadang'!$K${st}")
S.freeze_panes = "C6"
page(S)

# ======================================================================= CAPEX
C = wb.create_sheet("CAPEX")
title(C, "CAPEX Pemeliharaan (pengadaan sekali, tahun 1)",
      "Harga = estimasi harga pasar daring (marketplace) Oktober 2026 untuk barang setara; cek ulang saat pengadaan.",
      "Total = jumlah x harga x (1 + eskalasi)^(tahun ke - 1), supaya pengadaan yang digeser ke tahun lain ikut naik harga.")
header(C, 5, ["No", "Kategori", "Uraian", "Spesifikasi (setara)", "Fungsi dalam pemeliharaan", "Satuan", "Jumlah",
              "Harga satuan (Rp)", "Total (Rp)", "Tahun ke-", "Sumber harga"],
       [5, 24, 30, 30, 46, 8, 8, 14, 14, 8, 22])
K1, K2, K3, K4 = ("Peralatan kerja & alat ukur", "Infrastruktur perawatan", "Stok awal suku cadang kritis",
                  "Pelatihan & dokumentasi awal")
MP = "Estimasi marketplace Okt 2026"
capex = [
    (K1, "Laptop diagnostik", "Core i3/Ryzen 3, RAM 8 GB, SSD 256 GB", "Unggah firmware ESP32, baca log serial, simpan data riwayat perawatan (CMMS sederhana)", "unit", 1, 5500000, MP),
    (K1, "Osiloskop digital 2 kanal", "100 MHz (setara Hantek DSO2C10)", "Cek sinyal PWM motor, pulsa encoder, dan keluaran sensor IR saat diagnosis", "unit", 1, 2500000, MP),
    (K1, "Multimeter digital acuan", "True RMS (setara UNI-T UT61E+)", "Ukur tegangan sel & arus; jadi acuan verifikasi alat ukur lain (dikalibrasi tiap tahun)", "unit", 1, 850000, MP),
    (K1, "Multimeter digital lapangan", "Setara UNI-T UT33D+", "Pemeriksaan cepat saat PM mingguan", "unit", 1, 250000, MP),
    (K1, "Solder station suhu terkendali", "Setara Quicko T12", "Ganti konektor, header modul, kabel motor", "unit", 1, 850000, MP),
    (K1, "Hot air rework station", "Setara 858D", "Ganti sensor ITR20001 & komponen SMD di PCB sensor", "unit", 1, 750000, MP),
    (K1, "Catu daya DC laboratorium", "30 V / 5 A", "Uji modul & motor di meja tanpa baterai", "unit", 1, 900000, MP),
    (K1, "Penganalisis & pengisi baterai 18650", "Setara LiitoKala Lii-500 (ukur kapasitas & IR)", "Uji kapasitas dan tahanan dalam sel saat PM bulanan, menentukan sel yang diganti", "unit", 1, 450000, MP),
    (K1, "Tachometer optik digital", "Non-kontak", "Ukur RPM roda untuk mendeteksi motor yang mulai aus", "unit", 1, 250000, MP),
    (K1, "Jangka sorong digital", "150 mm", "Ukur keausan ban & dimensi suku cadang", "unit", 1, 200000, MP),
    (K1, "Set perkakas elektronik", "Obeng presisi, tang, pinset, crimper JST, penyedot timah", "Bongkar pasang robot & pembuatan kabel", "set", 1, 750000, MP),
    (K1, "Matras ESD + gelang antistatis + kaca pembesar berlampu", "-", "Mencegah kerusakan modul karena listrik statis saat perbaikan", "set", 1, 600000, MP),
    (K1, "Kabel USB-C data + USB hub", "-", "Koneksi laptop ke robot", "set", 1, 250000, MP),
    (K2, "Lemari suku cadang + kotak sortir berlabel", "-", "Penyimpanan stok suku cadang dengan sistem label (FIFO)", "unit", 1, 1200000, MP),
    (K2, "Stasiun pengisian baterai", "Rak + 5 adaptor USB-C 5 V/3 A + terminal", "Tempat pengisian 5 robot yang tertib dan aman", "set", 1, 600000, MP),
    (K2, "APAR serbuk ABC 3 kg", "-", "Keselamatan kebakaran di area pengisian baterai Li-ion", "unit", 1, 450000, MP),
    (K2, "Kantong/kotak baterai tahan api", "-", "Penyimpanan baterai cadangan & baterai rusak sebelum dibuang", "unit", 2, 150000, MP),
    (K2, "Lintasan uji & kalibrasi 4x4", "Papan + pita garis 19 mm, sama dengan lintasan PjBL", "Uji fungsi & kalibrasi sensor setelah perbaikan, sebelum robot kembali ke lini", "set", 1, 600000, MP),
    (K2, "Jig uji sensor & encoder", "Rangka + PCB uji sederhana", "Uji PCB sensor & motor tanpa membongkar seluruh robot", "set", 1, 400000, MP),
    (K3, "Stok awal suku cadang kritis", "Rincian di sheet Suku Cadang", "Mengurangi downtime: robot langsung diperbaiki tanpa menunggu pengiriman", "paket", 1, "=biaya_stok_awal", "Sheet Suku Cadang"),
    (K4, "Pelatihan awal teknisi & operator", "2 hari, oleh tim perancang", "Teknisi paham perawatan, keselamatan baterai Li-ion, unggah firmware; operator paham pemeriksaan awal shift", "paket", 1, 1500000, "Estimasi kelompok"),
    (K4, "Printer label aset / kode QR", "Setara Brother PT-E110", "Label aset tiap robot & suku cadang, terhubung ke kartu riwayat perawatan", "unit", 1, 500000, MP),
    (K4, "Cetak SOP, checklist PM & kartu riwayat", "Laminasi", "Dokumen kerja di stasiun perawatan", "paket", 1, 150000, "Estimasi kelompok"),
]
c0 = 6
for i, (kat, ur, spec, fungsi, sat, jml, harga, src) in enumerate(capex):
    r = c0 + i
    put(C, f"A{r}", i + 1, align="center")
    put(C, f"B{r}", kat, wrap=True)
    put(C, f"C{r}", ur, wrap=True)
    put(C, f"D{r}", spec, wrap=True)
    put(C, f"E{r}", fungsi, wrap=True)
    put(C, f"F{r}", sat)
    put(C, f"G{r}", jml, color=BLUE, fmt=NUM)
    put(C, f"H{r}", harga, color=GREEN if isinstance(harga, str) else BLUE, fmt=RP)
    put(C, f"I{r}", f"=G{r}*H{r}*(1+eskalasi)^(J{r}-1)", fmt=RP)
    put(C, f"J{r}", 1, color=BLUE, fmt=NUM, align="center")
    put(C, f"K{r}", src, wrap=True)
c1 = c0 + len(capex) - 1
ct = c1 + 2
put(C, f"C{ct}", "Total CAPEX", bold=True, fill=F_TOT)
put(C, f"I{ct}", f"=SUM(I{c0}:I{c1})", bold=True, fmt=RP, fill=F_TOT)
C.freeze_panes = "C6"
page(C)

# ======================================================================= OPEX
O = wb.create_sheet("OPEX")
title(O, "OPEX Pemeliharaan per Tahun (harga naik sesuai eskalasi)",
      "Biaya tahun t = volume x harga tahun 1 x (1 + eskalasi)^(t - 1), hanya pada tahun kelipatan interval.",
      "Volume suku cadang, jam preventif, dan jam penggantian diambil dari sheet Suku Cadang dan Jadwal PM.")
cols = ["No", "Kategori", "Uraian", "Dasar perhitungan", "Satuan", "Volume per tahun", "Harga satuan th-1 (Rp)",
        "Interval (tahun)", "Tahun 1", "Tahun 2", "Tahun 3", "Tahun 4", "Tahun 5", "Total 5 tahun (Rp)"]
header(O, 6, cols, [5, 22, 34, 40, 9, 10, 13, 9, 13, 13, 13, 13, 13, 15])
for k in range(5):            # baris bantu: indeks tahun & tahun kalender
    col = get_column_letter(9 + k)
    put(O, f"{col}5", k + 1, color=GREY, fmt="0", align="center", border=False)
    put(O, f"{col}7", f"=tahun_awal+{col}5-1", color=BLACK, fmt="0", align="center", fill=F_TOT)
put(O, "H5", "indeks tahun", color=GREY, border=False, size=8)
put(O, "H7", "tahun anggaran", fill=F_TOT, size=8)
O1, O2, O3, O4, O5, O6 = ("Suku cadang penggantian", "Bahan habis pakai", "Tenaga kerja pemeliharaan",
                          "Kalibrasi alat ukur", "Pelatihan penyegaran", "Listrik pengisian baterai")
opex = [
    (O1, "Suku cadang penggantian", "Kebutuhan per tahun dari sheet Suku Cadang", "paket", 1, "=biaya_sc_tahun", 1),
    (O2, "Timah solder + flux", "1 paket per triwulan", "paket", 4, 25000, 1),
    (O2, "Isopropyl alcohol 99% 1 L", "Pembersih sensor IR saat PM mingguan", "botol", 2, 90000, 1),
    (O2, "Cotton bud + kain mikrofiber", "1 paket per triwulan", "paket", 4, 40000, 1),
    (O2, "Kabel AWG26, heat shrink, cable tie", "Perbaikan kabel motor & konektor", "paket", 2, 75000, 1),
    (O2, "Baut, mur, spacer M3", "Pengganti baut aus/hilang", "paket", 4, 15000, 1),
    (O2, "Mata solder pengganti", "Aus pemakaian", "pcs", 2, 60000, 1),
    (O2, "Contact cleaner spray", "Pembersih konektor & saklar", "kaleng", 2, 65000, 1),
    (O2, "Pita garis lantai 19 mm (roll 33 m)", "Panjang lintasan x bagian yang diganti / 33 m per roll", "roll",
     "=ROUNDUP(panjang_lintasan*ganti_lintasan/33,0)", 180000, 1),
    (O3, "Pemeliharaan preventif oleh teknisi", "Jam teknisi dari sheet Jadwal PM", "jam", "=jam_pm", "=tarif_teknisi", 1),
    (O3, "Penggantian komponen (terencana + korektif)", "Jumlah penggantian x MTTR dari sheet Suku Cadang", "jam",
     "=jam_ganti", "=tarif_teknisi", 1),
    (O4, "Kalibrasi multimeter acuan (lab terakreditasi KAN)", "1x per tahun; alat lain diverifikasi internal terhadap multimeter acuan", "kali", 1, 400000, 1),
    (O5, "Pelatihan penyegaran teknisi & operator", "1 hari per tahun, materi + konsumsi", "paket", 1, 500000, 1),
    (O6, "Listrik pengisian baterai", "Jumlah pengisian x energi per pengisian / 1000", "kWh", "=n_cas*wh_cas/1000",
     "=tarif_listrik", 1),
]
o0 = 8
for i, (kat, ur, dasar, sat, vol, harga, itv) in enumerate(opex):
    r = o0 + i
    put(O, f"A{r}", i + 1, align="center")
    put(O, f"B{r}", kat, wrap=True)
    put(O, f"C{r}", ur, wrap=True)
    put(O, f"D{r}", dasar, wrap=True)
    put(O, f"E{r}", sat)
    vf = isinstance(vol, str)
    put(O, f"F{r}", vol, color=(GREEN if ("jam_" in str(vol) or "n_cas" in str(vol)) else BLACK) if vf else BLUE, fmt=NUM1)
    hf = isinstance(harga, str)
    put(O, f"G{r}", harga, color=GREEN if hf else BLUE, fmt=RP)
    put(O, f"H{r}", itv, color=BLUE, fmt=NUM, align="center")
    for k in range(5):
        col = get_column_letter(9 + k)
        put(O, f"{col}{r}", f"=IF(MOD({col}$5,$H{r})=0,$F{r}*$G{r}*(1+eskalasi)^({col}$5-1),0)", fmt=RP)
    put(O, f"N{r}", f"=SUM(I{r}:M{r})", fmt=RP)
o1 = o0 + len(opex) - 1
osub, okon, otot = o1 + 1, o1 + 2, o1 + 3
put(O, f"C{osub}", "Subtotal sebelum kontingensi", bold=True, fill=F_TOT)
put(O, f"B{okon}", "Kontingensi", fill=F_TOT)
put(O, f"C{okon}", "Kontingensi OPEX", fill=F_TOT)
put(O, f"D{okon}", "Persentase kontingensi x subtotal", fill=F_TOT)
put(O, f"C{otot}", "Total OPEX", bold=True, fill=F_TOT)
for k in range(6):
    col = get_column_letter(9 + k)
    put(O, f"{col}{osub}", f"=SUM({col}{o0}:{col}{o1})", bold=True, fmt=RP, fill=F_TOT)
    put(O, f"{col}{okon}", f"=kontingensi*{col}{osub}" if k < 5 else f"=SUM(I{okon}:M{okon})", fmt=RP, fill=F_TOT)
    put(O, f"{col}{otot}", f"={col}{osub}+{col}{okon}", bold=True, fmt=RP, fill=F_TOT)
O.freeze_panes = "D8"
page(O)

# ======================================================================= RINGKASAN
R = wb.create_sheet("Ringkasan", 0)
put(R, "A1", "RENCANA ANGGARAN BIAYA (RAB) CAPEX & OPEX PEMELIHARAAN - 5 TAHUN", bold=True, size=14, border=False)
put(R, "A2", "Sistem: armada AGV mini line tracer, 5 unit (4 beroperasi + 1 cadangan), hasil PjBL P01 Kelompok 5",
    border=False)
put(R, "A3", "Disusun oleh: Unit Maintenance  |  Diajukan kepada: Kepala Plant & Bagian Keuangan  |  "
             "Tugas Merancang RAB - Teknik Perawatan dan Perbaikan (B)", color=GREY, border=False, italic=True)
put(R, "A4", "Harga dalam Rupiah. Angka tahun 2-5 sudah termasuk eskalasi harga. Detail perhitungan di sheet "
             "Asumsi, Jadwal PM, Suku Cadang, CAPEX, dan OPEX.", color=GREY, border=False, italic=True)
header(R, 6, ["No", "Pos anggaran", "Tahun 1", "Tahun 2", "Tahun 3", "Tahun 4", "Tahun 5", "Total 5 tahun",
              "% dari total"], [6, 46, 15, 15, 15, 15, 15, 17, 11])
for k in range(5):
    col = get_column_letter(3 + k)
    put(R, f"{col}7", f"=tahun_awal+{k}", fmt="0", align="center", fill=F_TOT)
put(R, "B7", "Tahun anggaran", fill=F_TOT)
put(R, "A7", None, fill=F_TOT)
put(R, "H7", None, fill=F_TOT)
put(R, "I7", None, fill=F_TOT)
row = 8


def sec(label):
    global row
    put(R, f"A{row}", label[0], bold=True, fill=F_SEC)
    put(R, f"B{row}", label[1], bold=True, fill=F_SEC)
    for col in "CDEFGHI":
        put(R, f"{col}{row}", None, fill=F_SEC)
    row += 1


lines = {}
sec(("A", "CAPEX (pengadaan alat & stok awal)"))
capex_rows = []
for i, kat in enumerate((K1, K2, K3, K4)):
    put(R, f"A{row}", f"A.{i + 1}", align="center")
    put(R, f"B{row}", kat)
    for k in range(5):
        col = get_column_letter(3 + k)
        put(R, f"{col}{row}", f'=SUMIFS(CAPEX!$I${c0}:$I${c1},CAPEX!$B${c0}:$B${c1},$B{row},CAPEX!$J${c0}:$J${c1},{k + 1})',
            color=GREEN, fmt=RP)
    put(R, f"H{row}", f"=SUM(C{row}:G{row})", fmt=RP)
    capex_rows.append(row)
    row += 1
cap_sub = row
put(R, f"B{row}", "Subtotal CAPEX", bold=True, fill=F_TOT)
put(R, f"A{row}", None, fill=F_TOT)
for k in range(6):
    col = get_column_letter(3 + k)
    put(R, f"{col}{row}", f"=SUM({col}{capex_rows[0]}:{col}{capex_rows[-1]})", bold=True, fmt=RP, fill=F_TOT)
row += 1
sec(("B", "OPEX (biaya pemeliharaan rutin)"))
opex_rows = []
for i, kat in enumerate((O1, O2, O3, O4, O5, O6)):
    put(R, f"A{row}", f"B.{i + 1}", align="center")
    put(R, f"B{row}", kat)
    for k in range(5):
        col = get_column_letter(3 + k)
        ocol = get_column_letter(9 + k)
        put(R, f"{col}{row}", f'=SUMIFS(OPEX!{ocol}${o0}:{ocol}${o1},OPEX!$B${o0}:$B${o1},$B{row})', color=GREEN, fmt=RP)
    put(R, f"H{row}", f"=SUM(C{row}:G{row})", fmt=RP)
    opex_rows.append(row)
    row += 1
put(R, f"A{row}", "B.7", align="center")
put(R, f"B{row}", "Kontingensi (persentase di sheet Asumsi)")
for k in range(5):
    col = get_column_letter(3 + k)
    ocol = get_column_letter(9 + k)
    put(R, f"{col}{row}", f"=OPEX!{ocol}{okon}", color=GREEN, fmt=RP)
put(R, f"H{row}", f"=SUM(C{row}:G{row})", fmt=RP)
opex_rows.append(row)
row += 1
opx_sub = row
put(R, f"B{row}", "Subtotal OPEX", bold=True, fill=F_TOT)
put(R, f"A{row}", None, fill=F_TOT)
for k in range(6):
    col = get_column_letter(3 + k)
    put(R, f"{col}{row}", f"=SUM({col}{opex_rows[0]}:{col}{opex_rows[-1]})", bold=True, fmt=RP, fill=F_TOT)
row += 1
tot = row
put(R, f"A{row}", None, fill=F_HEAD)
put(R, f"B{row}", "TOTAL RAB (CAPEX + OPEX)", color=WHITE, bold=True, fill=F_HEAD)
for k in range(6):
    col = get_column_letter(3 + k)
    put(R, f"{col}{row}", f"={col}{cap_sub}+{col}{opx_sub}", color=WHITE, bold=True, fmt=RP, fill=F_HEAD)
put(R, f"I{row}", None, fill=F_HEAD)
# kolom persentase
for rr in capex_rows + [cap_sub] + opex_rows + [opx_sub, tot]:
    put(R, f"I{rr}", f"=IF($H${tot}=0,0,H{rr}/$H${tot})", fmt=PCT,
        color=WHITE if rr == tot else BLACK, bold=rr in (cap_sub, opx_sub, tot),
        fill=F_HEAD if rr == tot else (F_TOT if rr in (cap_sub, opx_sub) else None))
row += 2
put(R, f"B{row}", "Kesesuaian dengan pagu", bold=True, border=False)
row += 1
chk = [
    ("Pagu anggaran 5 tahun", "=pagu", RP, GREEN),
    ("Total RAB 5 tahun", f"=H{tot}", RP, BLACK),
    ("Sisa pagu", f"=pagu-H{tot}", RP, BLACK),
    ("Persentase pagu terpakai", f"=IF(pagu=0,0,H{tot}/pagu)", PCT, BLACK),
    ("Status", f'=IF(H{tot}<=pagu,"SESUAI - di bawah pagu","MELEBIHI PAGU - kurangi pos")', None, BLACK),
]
for lab, f, fmt, colr in chk:
    put(R, f"B{row}", lab)
    put(R, f"C{row}", f, fmt=fmt, color=colr, bold=lab == "Status")
    row += 1
row += 1
put(R, f"B{row}", "Indikator pemeliharaan", bold=True, border=False)
row += 1
ind = [
    ("Rata-rata biaya pemeliharaan per robot per tahun (Rp)", f"=H{tot}/n_robot/5", RP),
    ("Rata-rata biaya pemeliharaan per jam jalan robot (Rp/jam)", f"=H{tot}/(jam_jalan*5)", RP),
    ("Jam teknisi preventif per tahun", "=jam_pm", NUM1),
    ("Jam teknisi penggantian komponen per tahun", "=jam_ganti", NUM1),
    ("Porsi jam preventif dari total jam teknisi", "=IF(jam_pm+jam_ganti=0,0,jam_pm/(jam_pm+jam_ganti))", PCT),
    ("Jumlah penggantian komponen per tahun (kejadian)", "=n_ganti", NUM),
    ("Biaya teknisi per jam (Rp)", "=tarif_teknisi", RP),
]
for lab, f, fmt in ind:
    put(R, f"B{row}", lab, wrap=True)
    put(R, f"C{row}", f, fmt=fmt, color=GREEN)
    row += 1
row += 1
put(R, f"B{row}", "Catatan", bold=True, border=False)
row += 1
notes = [
    "1. Harga robot tidak termasuk: robot sudah terpasang dan beroperasi; RAB ini hanya biaya pemeliharaannya.",
    "2. CAPEX dibeli di tahun 1. Stok awal suku cadang kritis disimpan supaya robot rusak bisa langsung diperbaiki.",
    "3. Kebutuhan suku cadang dihitung dari jam jalan (motor, ban, caster), siklus baterai, dan laju kerusakan modul.",
    "4. Strategi: preventif terjadwal (sheet Jadwal PM) + penggantian berbasis kondisi dari hasil ukur baterai & RPM.",
    "5. Ubah asumsi di sheet Asumsi (sel kuning) bila data riwayat tahun pertama menunjukkan angka berbeda.",
]
for n in notes:
    put(R, f"B{row}", n, border=False)
    row += 1
R.freeze_panes = "C8"
page(R)

wb.calculation = CalcProperties(fullCalcOnLoad=True)
wb.save(OUT)
print("saved", OUT)
